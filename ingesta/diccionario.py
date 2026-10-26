"""Actualiza el Excel del diccionario de la entrega 1 con lo que de verdad está hoy en el lago (entrega 2).

    .venv/bin/python ingesta/diccionario.py

Parte de la plantilla entregada (docs/entrega1/Diccionario_de_datos_TallerDatos_completo.xlsx), que no se toca, y escribe
docs/Diccionario_de_datos_TallerDatos_entrega2.xlsx:

- 1. Fuentes: la licencia copiada de la fuente (ingesta/licencias.py; «no declara» si la fuente no dice nada), la fecha real de
  prueba y el periodo del catálogo, y lo que trae la fuente sobre personas. Se agregan las fuentes que no estaban (Open-Meteo).
- 2. Diccionario: se quitan los campos que ya no existen en el lago y se agregan los de la entrega 2.
- 3. Bitácora IA: se agregan los encargos de la entrega 2, con lo que la IA hizo mal y cómo se corrigió.
- Protección y Cambios entrega 2: hojas nuevas.

Los IDs (F001…) son los mismos de la entrega 1, de la app (sección Fuentes) y del README.
"""

import copy
import json
import re
import sys

import openpyxl

from lago import DIR_LAGO, RAIZ, hoy

PLANTILLA = RAIZ / 'docs' / 'entrega1' / 'Diccionario_de_datos_TallerDatos_completo.xlsx'
SALIDA = RAIZ / 'docs' / 'Diccionario_de_datos_TallerDatos_entrega2.xlsx'
FILA_EJEMPLO = 4  # la fila gris de la plantilla del profesor (F01, Wikipedia): no se toca

# Columnas de «1. Fuentes» (1 = A).
C_ID, C_NOMBRE, C_ENTIDAD, C_URL, C_LICENCIA, C_PERSONAS, C_PERIODO, C_PROBADA, C_ESTADO, C_OBS = 1, 2, 3, 5, 8, 9, 11, 13, 14, 16

# Lo que trae el archivo de origen de las fuentes con datos de personas y lo que llega al lago. En la entrega 1 la columna
# «¿Trae datos de personas?» del SISC decía «Solo conteos por zona» mientras su observación decía que cada evento trae sexo,
# edad y coordenadas: la columna describía el lago, no la fuente. La fuente SÍ los trae; el lago NO.
PERSONAS_EN_LA_FUENTE = {
    'sisc-': 'Personas identificables',
    'medata-victimas-viales': 'Personas identificables',
    'secop2-distrito': 'Personas identificables',
    'gdelt-titulares': 'Texto libre sin revisar',
}

CAMPOS_CRUDOS = [
    ('SISC (F020–F027)', 'fecha_hecho, cantidad', 'Sí: fecha y hora del hecho', 'Solo el año, en conteos'),
    ('SISC (F020–F027)', 'latitud, longitud', 'Sí: coordenadas de cada evento', 'No: sirven para ubicar el evento en su comuna y su barrio, y se descartan'),
    ('SISC (F020–F027)', 'codigo_comuna, codigo_barrio, nombre_barrio', 'Sí', 'Agregado: conteos por comuna (un año) y por barrio (tres años, solo si son 5 o más)'),
    ('SISC (F020–F027)', 'sexo, edad, estado_civil, ocupacion, discapacidad, grupo_especial, nivel_academico, parentesco',
     'Sí: datos de la víctima', 'No: nunca salen del equipo que procesa el archivo (datos/crudos/ está fuera de git y del despliegue)'),
    ('Víctimas viales (F065)', 'Sexo, Edad, Condicion, Direccion_incidente, Radicado, Fecha_incidente, Hora_incidente',
     'Sí: datos de la víctima y del lugar exacto', 'No'),
    ('Víctimas viales (F065)', 'Latitud, Longitud, Año, Gravedad_victima', 'Sí',
     'Mapa de calor: el centro de una celda de ~110 m y el año, sin la gravedad. Conteos por comuna.'),
    ('Policía Nacional (F013, F018, F019)', 'sexo', 'Sí: sexo de la víctima por evento, sin lugar más fino que el municipio', 'No: solo se piden sumas por año'),
    ('SECOP II (F040)', 'proveedor, tipodocproveedor, documento_proveedor', 'Sí: el proveedor puede ser una persona natural',
     'Solo el nombre de los 10 contratos de mayor valor, que son de empresas; nunca el documento'),
    ('Prensa vía GDELT (F092)', 'título, medio, fecha, URL', 'Sí: un titular puede nombrar a una persona',
     'Con filtro: no se publican los titulares que cuentan una muerte, una captura, un ataque o una desaparición'),
]

# Encargos de la entrega 2. La columna de lo que se corrigió es la que importa: cada fila dice qué hizo mal la IA.
BITACORA = [
    ('Claude Code', 'Revisar la contradicción del SISC que señaló el profesor',
     'Abrió la cabecera del CSV de homicidios: trae sexo, edad, estado civil, ocupación, discapacidad, grupo especial y coordenadas de cada víctima',
     'curl a los primeros bytes de 4 archivos del SISC en MEData', 'Completo', 'F020–F027',
     'En la entrega 1 la IA llenó «¿Trae datos de personas?» pensando en lo que guarda el lago y no en lo que trae la fuente. Ninguno de los dos lo revisó.'),
    ('Claude Code', 'Explicar la rentabilidad de 18 % en estrato 1',
     'Encontró que el arriendo de estrato 1 eran 31 ofertas (29 apartamentos) y la venta 366 (303 casas a la mitad del precio por m²)',
     'Se repitió la consulta a la OIME separando casas y apartamentos', 'Completo', 'F008',
     'Error de la IA en la entrega 1: dividió medianas de inmuebles distintos. Ahora compara apartamentos con apartamentos y exige 30 ofertas de cada lado.'),
    ('Claude Code', 'Corregir la tasa del SISC',
     'La tasa de la entrega 1 sumaba 3 años (el último cortado el 29 de noviembre de 2023) y los dividía entre la población del DAP de 2026, 10,3 % más alta que la del DANE',
     'Se leyó el código de ingesta y la nota del indicador', 'Completo', 'F020–F027, F001, F003',
     'La IA había escrito «No es una tasa anual» en la nota en vez de arreglarla. Ahora: un año completo sobre la población de ese mismo año.'),
    ('Claude Code', 'Corregir el verde por habitante',
     'El inventario es de 2019 y se dividía entre la población de 2026; la nota lo advertía pero la cifra seguía mezclando años',
     'Se comparó la vigencia de numerador y denominador', 'Completo', 'F078',
     'Una advertencia en la nota no arregla una cifra: ahora inventario y población son de 2019.'),
    ('Claude Code', 'Leer la licencia de las 93 fuentes',
     'Automatizó la lectura en datos.gov.co y MEData; para los servidores ArcGIS propuso leer copyrightText',
     'Se compararon con las fichas: copyrightText es un crédito, no una licencia', 'Parcial', 'F001–F093',
     'Se registró «no declara» con el crédito en observaciones. 10 fuentes del portal del Área Metropolitana quedaron sin verificar porque respondía 502.'),
    ('Claude Code', 'Numerar las fuentes con un código',
     'Las numeró F01–F93 por su cuenta, sin mirar el Excel de la entrega 1; desde F057 no coincidía con nuestros IDs',
     'Se cruzó cada URL del catálogo con la hoja 1. Fuentes', 'Completo', 'F057–F093',
     'Los IDs los pone el grupo, no la IA: se volvió a los de la entrega 1 y Open-Meteo, que no estaba, quedó como F093.'),
    ('Claude Code', 'Revisar si hay datos de personas a la vista',
     'Encontró titulares de prensa que nombran víctimas («Falleció joven de 20 años…», capturas con alias) y barrios con 1 a 4 homicidios en el mapa',
     'Se leyeron 40 titulares al azar y se contaron los barrios con menos de 5 casos (161 de 239 en homicidios)', 'Completo', 'F092, F020–F027',
     'Se agregó un filtro de titulares sobre personas y el mapa por barrio ya no publica conteos menores de 5.'),
    ('Claude Code', 'Probar Streamlit con la misma tabla',
     'La app funcionó; st.bar_chart ordenó las barras por nombre y no por valor, y el mapa por comuna pidió instalar pydeck',
     'Se corrió la app y se tomó una captura (exploracion/streamlit)', 'Parcial', 'F020',
     'El entorno virtual falló con Python 3.14 (sin ensurepip) y hubo que usar 3.12.'),
    ('Claude Code', 'Probar Observable Framework con la misma tabla',
     'El mapa salió como un rectángulo lleno y las barras no se ordenaron al primer intento',
     'Capturas antes y después (exploracion/observable)', 'Parcial', 'F020',
     'Los polígonos de ArcGIS vienen con el giro contrario al que espera d3; se resolvió con la proyección reflect-y.'),
    ('Claude Code', 'Probar kepler.gl con la misma tabla',
     'Abrió la demo y subió el GeoJSON, pero el clic en «Add Data» no se completó en un Chrome sin ventana; el grupo lo terminó a mano',
     'Capturas de cada paso (exploracion/kepler) y la prueba a mano de Juan Jose', 'Completo', 'F020',
     'La IA no logró automatizar el último paso. Ya en el navegador, kepler eligió una escala por cuantiles que no muestra magnitudes y no lleva el ID de la fuente.'),
    ('Claude Code', 'Errores de la entrega 1 que no se habían registrado',
     'La afluencia del Metro salió en 1,99 millones por día: sumaba la columna de total dos veces (son 997.000)',
     'Se comparó con el total que publica el Metro', 'Completo', 'F055',
     'Se registra ahora: en la entrega 1 la bitácora decía que todo había salido bien.'),
    ('Claude Code', 'Errores de la entrega 1 que no se habían registrado',
     'Al dividir el mapa 3D en archivos dejó una función sin importar: tocar el mapa no abría ninguna ficha. Además mostraba «-999 dB» (el código de «sin dato» del SIATA) como una lectura real',
     'Prueba con clics en un Chrome sin ventana (npm run qa)', 'Completo', 'F084',
     'El lint y las pruebas unitarias no lo veían: solo hacer clic lo encontró.'),
]


def cargar():
    temas = {r.stem: json.loads(r.read_text(encoding='utf-8')) for r in sorted(DIR_LAGO.glob('*.json'))
             if r.name not in ('indice.json', 'catalogo.json')}
    catalogo = json.loads((DIR_LAGO / 'catalogo.json').read_text(encoding='utf-8'))
    return temas, catalogo


def existe(temas, archivo, campo):
    """¿El campo de la hoja 2 («cifras.homicidios.valor», «listas.ecas[].capacidad») está en el lago?"""
    tema = temas.get(re.sub(r'^lago/|\.json$', '', archivo or ''))
    if not tema or not campo:
        return False
    partes = campo.replace('[]', '').split('.')
    if partes[0] == 'probado':
        return 'probado' in tema
    if partes[0] == 'territorios':
        partes = ['listas'] + partes
    if partes[0] in ('cifras', 'series'):
        coleccion = tema[partes[0]]
        if partes[1] == '<clave>':
            return bool(coleccion)
        return partes[1] in coleccion and (len(partes) < 3 or partes[2] in coleccion[partes[1]])
    if partes[0] == 'listas':
        lista = tema['listas'].get(partes[1])
        if not lista:
            return False
        return len(partes) < 3 or any(partes[2] in fila for fila in lista if isinstance(fila, dict))
    return False


def copiar_estilo(hoja, origen, destino):
    for col in range(1, hoja.max_column + 1):
        de, a = hoja.cell(origen, col), hoja.cell(destino, col)
        if de.has_style:
            a._style = copy.copy(de._style)


def actualizar_fuentes(hoja, catalogo, cambios):
    filas = {hoja.cell(f, C_ID).value: f for f in range(FILA_EJEMPLO + 1, hoja.max_row + 1) if hoja.cell(f, C_ID).value}
    for d in catalogo['datasets']:
        fila = filas.get(d['codigo'])
        if fila is None:
            fila = hoja.max_row + 1
            copiar_estilo(hoja, fila - 1, fila)
            for col, valor in ((C_ID, d['codigo']), (C_NOMBRE, d['nombre']), (C_ENTIDAD, d['entidad']), (4, 'Entidad privada'),
                               (C_URL, d['url']), (6, 'JSON'), (7, 'API sin llave'), (10, 'Medellín'), (12, 'En vivo'),
                               (15, 'Búsqueda del grupo con IA')):
                hoja.cell(fila, col, valor)
            cambios.append((d['codigo'], '1. Fuentes', 'Fuente agregada: estaba en la app y no en el Excel de la entrega 1'))
        licencia = d['licencia']
        detalle = ''
        if licencia.startswith('No declarada'):
            # La plantilla pide «no declara» cuando la fuente no dice nada; lo que sí dice (un crédito) va a observaciones.
            detalle = licencia.split('. ', 1)[1] if '. ' in licencia else ''
            licencia = 'no declara'
        anterior = hoja.cell(fila, C_LICENCIA).value
        if anterior != licencia:
            cambios.append((d['codigo'], '1. Fuentes', f'Licencia: «{anterior}» → «{licencia[:120]}»'))
        hoja.cell(fila, C_LICENCIA, licencia)
        hoja.cell(fila, C_PERIODO, d['vigencia'])
        hoja.cell(fila, C_PROBADA, d['probado'])
        hoja.cell(fila, C_ESTADO, 'Candidata' if d['estado'] == 'candidato' else 'Integrada al lago')
        if not hoja.cell(fila, C_PERSONAS).value:
            hoja.cell(fila, C_PERSONAS, 'No')
        personas = next((v for k, v in PERSONAS_EN_LA_FUENTE.items() if d['id'].startswith(k)), None)
        if personas and hoja.cell(fila, C_PERSONAS).value != personas:
            cambios.append((d['codigo'], '1. Fuentes', f'¿Trae datos de personas?: «{hoja.cell(fila, C_PERSONAS).value}» → «{personas}» '
                                                      '(la columna describe la fuente, no el lago)'))
            hoja.cell(fila, C_PERSONAS, personas)
        obs = re.sub(r' · Entrega 2:.*$', '', str(hoja.cell(fila, C_OBS).value or ''))
        nota = f' · Entrega 2: {d["proteccion"].rstrip(".")}. Licencia leída el {d["licencia_probada"]} en {d["licencia_url"]}.'
        if detalle:
            nota += f' {detalle}.'
        hoja.cell(fila, C_OBS, (obs + nota).strip(' ·'))
    return hoja.max_row


def actualizar_diccionario(hoja, temas, cambios):
    for fila in range(hoja.max_row, FILA_EJEMPLO, -1):
        archivo, campo, codigo = hoja.cell(fila, 2).value, hoja.cell(fila, 3).value, hoja.cell(fila, 1).value
        if campo and not existe(temas, archivo, campo):
            cambios.append((codigo, '2. Diccionario', f'Campo quitado: {archivo} · {campo} ya no está en el lago'))
            hoja.delete_rows(fila)
    nuevas = [
        ('F020', 'lago/seguridad.json', 'listas.territorios[].valores.sisc_homicidio_tasa', 'Homicidio (SISC)',
         'Homicidios de un año completo por cada 10.000 habitantes de ese mismo año', 'Decimal', 'por 10.000 hab.', '25,10', 'Sí',
         '≥ 0; nulo si falta la población', 'Pública',
         'Casos del último año completo ÷ población del año (total DANE repartido según el DAP) × 10.000. Entrega 2: antes sumaba 3 años.',
         '¿Cuáles son las zonas con más homicidios?'),
        ('F006', 'lago/demografia.json', 'listas.territorios[].valores.natalidad_tasa', 'Natalidad',
         'Nacidos vivos por cada 100.000 habitantes de la comuna, calculado por la fuente', 'Decimal', 'por 100.000 hab.', '1086,8', 'Sí',
         '≥ 0', 'Pública', 'Se toma tal cual; la Secretaría de Salud la calcula con las proyecciones del DAP.',
         '¿Cuál es la comuna con mayor natalidad?'),
        ('F078', 'lago/territorio.json', 'listas.territorios[].valores.m2_verde_hab', 'area_ha',
         'Metros cuadrados de espacio verde urbano por habitante', 'Decimal', 'm²/hab.', '57,8', 'Sí', '≥ 0', 'Pública',
         'Área del inventario 2019 ÷ población de 2019. Entrega 2: antes se dividía entre la población de 2026.',
         '¿Cuánto espacio verde urbano hay por habitante en cada comuna?'),
        ('F001', 'lago/demografia.json', 'listas.poblacion_anual_por_territorio[].valores', 'total_2018 … total_2030 (DAP)',
         'Población de cada comuna año por año, para dividir cada cifra entre la población de su mismo año', 'Entero', 'habitantes',
         '141840', 'No', '> 0', 'Pública', 'Total DANE del año × proporción del territorio en la proyección del DAP de ese año (PROXY).',
         '¿Cuántos habitantes tenía cada comuna el año de la cifra?'),
        ('F001', 'lago/catalogo.json', 'datasets[].codigo · licencia · personas · proteccion', 'n/a',
         'ID fijo de la fuente, licencia copiada de la fuente, qué trae sobre personas y cómo se protege', 'Texto', 'n/a',
         'F020 · Atribución 4.0 Internacional (CC BY 4.0)', 'No', 'ID F000; protección: Abierto, solo agregados, No se publica', 'Pública',
         'ingesta/licencias.py lee la licencia en la fuente; ingesta/registro_fuentes.json la guarda.', '¿De dónde sale cada número y se puede usar?'),
        ('F008', 'lago/economia.json', 'cifras.<clave>.proxy', 'n/a',
         'Si el dato aproxima otra cosa que no se mide directo, de qué es aproximación', 'Texto', 'n/a',
         'el precio de cierre, que no es dato abierto', 'Sí', 'texto libre o vacío', 'Pública', 'Se asigna en el script de ingesta.',
         '¿Qué cifras son aproximaciones (PROXY)?'),
    ]
    for valores in nuevas:
        fila = hoja.max_row + 1
        copiar_estilo(hoja, fila - 1, fila)
        for col, v in enumerate(valores, 1):
            hoja.cell(fila, col, v)
        cambios.append((valores[0], '2. Diccionario', f'Campo agregado: {valores[1]} · {valores[2]}'))


def actualizar_bitacora(hoja):
    ultimo = max((int(v) for v in (hoja.cell(f, 1).value for f in range(FILA_EJEMPLO + 1, hoja.max_row + 1))
                  if str(v or '').isdigit()), default=0)
    for i, (herramienta, encargo, respuesta, comprobacion, resultado, fuente, aprendizaje) in enumerate(BITACORA, 1):
        fila = hoja.max_row + 1
        copiar_estilo(hoja, fila - 1, fila)
        for col, v in enumerate((str(ultimo + i), herramienta, f'Entrega 2 · {encargo}', respuesta, comprobacion, resultado, fuente,
                                 aprendizaje), 1):
            hoja.cell(fila, col, v)


def hoja_simple(libro, nombre, cabecera, filas, anchos):
    hoja = libro.create_sheet(nombre)
    hoja.append(cabecera)
    for f in filas:
        hoja.append(list(f))
    for i, ancho in enumerate(anchos):
        hoja.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = ancho
    for celda in hoja[1]:
        celda.font = openpyxl.styles.Font(bold=True)
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')


def main():
    temas, catalogo = cargar()
    libro = openpyxl.load_workbook(PLANTILLA)
    cambios = []
    ins = libro['Instrucciones']
    ins['B1'] = 'Taller de datos · Entrega 2 · El lago y su diccionario, actualizado'
    ins.cell(ins.max_row + 2, 2, f'Actualizado el {hoy().isoformat()} por ingesta/diccionario.py desde el lago. La hoja «Cambios entrega 2» '
                                 'lista cada cambio frente a la entrega 1, y la hoja «Protección» qué trae cada fuente sobre personas y qué llega al lago.')
    ultima_fuente = actualizar_fuentes(libro['1. Fuentes'], catalogo, cambios)
    actualizar_diccionario(libro['2. Diccionario'], temas, cambios)
    actualizar_bitacora(libro['3. Bitácora IA'])
    # Las fórmulas del resumen cubren las filas 5 a 96: se extienden si se agregaron fuentes.
    for fila in libro['Resumen'].iter_rows():
        for celda in fila:
            if isinstance(celda.value, str) and "'1. Fuentes'" in celda.value:
                celda.value = re.sub(r'(\$?[A-Z]+)96', rf'\g<1>{max(96, ultima_fuente)}', celda.value)
    hoja_simple(libro, 'Protección', ['Fuente', 'Campos del archivo de origen', '¿Los trae la fuente?', '¿Llegan al lago o a la app?'],
                CAMPOS_CRUDOS, [32, 60, 50, 80])
    hoja_simple(libro, 'Cambios entrega 2', ['ID', 'Hoja', 'Cambio'], cambios, [9, 16, 120])
    libro.save(SALIDA)
    print(f'  ✓ {SALIDA.relative_to(RAIZ)} · {len(catalogo["datasets"])} fuentes · {len(cambios)} cambios frente a la entrega 1')
    return 0


if __name__ == '__main__':
    sys.exit(main())
